from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Mapping, Sequence
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "17233C"
TEAL = "0D5F63"
TEAL_TEXT = "0D786F"
GOLD = "F2C45C"
INK = "102528"
MUTED = "496164"
WHITE = "FFFFFF"
PALE_TEAL = "D8F1E9"
PALE_GREEN = "EAF4EA"
PALE_CORAL = "F7E4DD"
LINE = "DDE4E1"


def build_schedule_workbook(
    rows: Sequence[Mapping[str, str]],
    *,
    slots: Sequence[str],
    timezone_name: str,
    as_of_date: date | None = None,
    horizon_days: int = 14,
) -> bytes:
    """Return a styled Excel schedule matrix built from authoritative tracker rows."""
    if not slots:
        raise ValueError("At least one publishing slot is required.")
    zone = ZoneInfo(timezone_name)
    start = as_of_date or datetime.now(zone).date()
    scheduled = [
        dict(row)
        for row in rows
        if str(row.get("publishing_date", "")).strip()
        or str(row.get("successfully_published", "")).casefold() == "yes"
    ]

    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "Schedule Matrix"
    availability = workbook.create_sheet("Slot Availability")
    _build_matrix_sheet(matrix, scheduled, slots, timezone_name)
    _build_availability_sheet(
        availability,
        scheduled,
        slots,
        timezone_name,
        start,
        horizon_days,
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _build_matrix_sheet(sheet, rows, slots, timezone_name: str) -> None:
    headers = [
        "S.No.",
        "YouTube video / Short",
        f"Publish date ({_timezone_label(timezone_name)})",
        "Status",
        *slots,
        "Catch-up / other",
    ]
    last_col = len(headers)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = TEAL_TEXT
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sheet.cell(2, 1, "Talkivana publishing schedule matrix")
    sheet.cell(2, 1).font = Font(name="Arial", size=16, bold=True, color=INK)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    sheet.cell(
        3,
        1,
        "A tick identifies the assigned publishing slot. Catch-up times appear in the final column.",
    )
    sheet.cell(3, 1).font = Font(name="Arial", size=10, italic=True, color=MUTED)
    _write_header(sheet, 5, headers, TEAL)

    for index, row in enumerate(rows, start=1):
        excel_row = index + 5
        publishing_date = str(row.get("publishing_date", ""))
        day, inferred_time = _split_local_datetime(publishing_date)
        slot = str(row.get("publish_slot_local", "")) or inferred_time
        fixed_slot = slot in slots
        values = [
            index,
            str(row.get("youtube_video_or_short_name", "")),
            day,
            str(row.get("publication_status", "")),
            *("✓" if candidate == slot else "" for candidate in slots),
            "" if fixed_slot else slot,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, column, value)
            cell.font = Font(name="Arial", size=10, color="26383A")
            cell.alignment = Alignment(
                horizontal="center" if column >= 5 else "left",
                vertical="center",
                wrap_text=column in (2, 4),
            )
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        for column in range(5, 5 + len(slots)):
            if sheet.cell(excel_row, column).value == "✓":
                sheet.cell(excel_row, column).fill = PatternFill("solid", fgColor=PALE_TEAL)
                sheet.cell(excel_row, column).font = Font(
                    name="Arial", size=12, bold=True, color=TEAL_TEXT
                )

    sheet.freeze_panes = "C6"
    sheet.auto_filter.ref = f"A5:{get_column_letter(last_col)}{max(5, 5 + len(rows))}"
    widths = [8, 45, 20, 34] + [11] * len(slots) + [17]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[5].height = 30


def _build_availability_sheet(
    sheet,
    rows,
    slots,
    timezone_name: str,
    start: date,
    horizon_days: int,
) -> None:
    headers = ["Date", "Day", *slots, "Open slots", "Earliest open"]
    last_col = len(headers)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = GOLD
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sheet.cell(2, 1, f"Publishing-slot availability · next {horizon_days} days")
    sheet.cell(2, 1).font = Font(name="Arial", size=16, bold=True, color=INK)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    sheet.cell(
        3,
        1,
        "Regenerated from the same scheduled-video registry checked before the automation assigns a time.",
    )
    sheet.cell(3, 1).font = Font(name="Arial", size=10, italic=True, color=MUTED)
    _write_header(sheet, 5, headers, NAVY)

    occupancy: dict[tuple[date, str], str] = {}
    for row in rows:
        day_text, inferred_time = _split_local_datetime(str(row.get("publishing_date", "")))
        if not day_text:
            continue
        try:
            day = date.fromisoformat(day_text)
        except ValueError:
            continue
        slot = str(row.get("publish_slot_local", "")) or inferred_time
        occupancy[(day, slot)] = str(row.get("youtube_video_or_short_name", ""))

    for offset in range(horizon_days):
        day = start + timedelta(days=offset)
        excel_row = 6 + offset
        values = [day, day.strftime("%a")]
        values.extend(occupancy.get((day, slot), "Available") for slot in slots)
        values.extend([None, None])
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, column, value)
            cell.font = Font(name="Arial", size=10, color="26383A")
            cell.alignment = Alignment(vertical="center", wrap_text=3 <= column < 3 + len(slots))
            cell.border = Border(bottom=Side(style="thin", color=LINE))
        sheet.cell(excel_row, 1).number_format = "yyyy-mm-dd"
        first_slot_col = 3
        last_slot_col = first_slot_col + len(slots) - 1
        slot_range = f"{get_column_letter(first_slot_col)}{excel_row}:{get_column_letter(last_slot_col)}{excel_row}"
        sheet.cell(excel_row, last_slot_col + 1, f'=COUNTIF({slot_range},"Available")')
        formula = '"Full"'
        for slot_index in range(len(slots) - 1, -1, -1):
            column = get_column_letter(first_slot_col + slot_index)
            formula = f'IF({column}{excel_row}="Available","{slots[slot_index]}",{formula})'
        sheet.cell(excel_row, last_slot_col + 2, f"={formula}")

    end_row = 5 + horizon_days
    slot_start = get_column_letter(3)
    slot_end = get_column_letter(2 + len(slots))
    cells = f"{slot_start}6:{slot_end}{end_row}"
    sheet.conditional_formatting.add(
        cells,
        FormulaRule(
            formula=[f'{slot_start}6="Available"'],
            fill=PatternFill("solid", fgColor=PALE_GREEN),
            font=Font(name="Arial", color="557557"),
        ),
    )
    sheet.conditional_formatting.add(
        cells,
        FormulaRule(
            formula=[f'AND({slot_start}6<>"",{slot_start}6<>"Available")'],
            fill=PatternFill("solid", fgColor=PALE_CORAL),
            font=Font(name="Arial", bold=True, color="8D4338"),
        ),
    )
    note_row = end_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=last_col)
    sheet.cell(
        note_row,
        1,
        "Manual workbook changes do not reschedule YouTube. Use the approved GitHub workflow so the manifest, matrix, tracker, and channel stay synchronized.",
    )
    sheet.cell(note_row, 1).fill = PatternFill("solid", fgColor="FFF3CD")
    sheet.cell(note_row, 1).font = Font(name="Arial", size=10, bold=True, color="654D14")
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[note_row].height = 34
    sheet.freeze_panes = "C6"
    widths = [14, 10] + [28] * len(slots) + [12, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[5].height = 30


def _write_header(sheet, row_number: int, headers: Sequence[str], fill_color: str) -> None:
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row_number, column, value)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color=WHITE))


def _split_local_datetime(value: str) -> tuple[str, str]:
    if not value:
        return "", ""
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return value[:10], value[11:16] if "T" in value else ""
    return parsed.date().isoformat(), parsed.strftime("%H:%M")


def _timezone_label(timezone_name: str) -> str:
    return "IST" if timezone_name == "Asia/Kolkata" else timezone_name

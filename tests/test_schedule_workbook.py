import unittest
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from language_video_automation.schedule_workbook import build_schedule_workbook


class ScheduleWorkbookTests(unittest.TestCase):
    def test_workbook_ticks_assigned_slot_and_marks_availability(self):
        content = build_schedule_workbook(
            [{
                "youtube_video_or_short_name": "Video 1: Test lesson",
                "publishing_date": "2026-10-26T18:00+05:30",
                "publish_slot_local": "18:00",
                "publication_status": "Scheduled; currently private",
                "successfully_published": "No",
            }],
            slots=["07:30", "12:30", "18:00", "22:30"],
            timezone_name="Asia/Kolkata",
            as_of_date=date(2026, 10, 26),
        )
        workbook = load_workbook(BytesIO(content), data_only=False)
        matrix = workbook["Schedule Matrix"]
        availability = workbook["Slot Availability"]
        self.assertEqual(matrix["G6"].value, "✓")
        self.assertIsNone(matrix["E6"].value)
        self.assertEqual(availability["E6"].value, "Video 1: Test lesson")
        self.assertEqual(availability["C6"].value, "Available")
        self.assertEqual(availability["G6"].value, '=COUNTIF(C6:F6,"Available")')
        self.assertEqual(availability["H6"].value, '=IF(C6="Available","07:30",IF(D6="Available","12:30",IF(E6="Available","18:00",IF(F6="Available","22:30","Full"))))')

    def test_catch_up_time_uses_other_column(self):
        content = build_schedule_workbook(
            [{
                "youtube_video_or_short_name": "Short 1: Catch-up",
                "publishing_date": "2026-10-26T16:50+05:30",
                "publish_slot_local": "16:50",
                "publication_status": "Scheduled; currently private",
                "successfully_published": "No",
            }],
            slots=["07:30", "12:30", "18:00", "22:30"],
            timezone_name="Asia/Kolkata",
            as_of_date=date(2026, 10, 26),
        )
        workbook = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(workbook["Schedule Matrix"]["I6"].value, "16:50")


if __name__ == "__main__":
    unittest.main()
